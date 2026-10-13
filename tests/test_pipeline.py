"""Real CLI/SQLite integration; only external news and AI endpoints are simulated."""
import json
import sqlite3
import subprocess
import sys
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def site():
    state = {'version': 1, 'fail_ai': False, 'ai_calls': 0, 'deny': False}

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def do_GET(self):
            base = f'http://127.0.0.1:{self.server.server_port}'
            if self.path == '/robots.txt':
                body = 'User-agent: *\nDisallow: /article/' if state['deny'] else 'User-agent: *\nAllow: /'
            elif self.path == '/feed':
                body = f'''<rss version="2.0"><channel><title>Test</title><link>{base}</link><description>Test feed</description>
                <item><title>우주 뉴스</title><link>{base}/article/1?utm_source=rss</link><pubDate>Wed, 07 Oct 2026 09:00:00 +0900</pubDate><category>우주</category></item>
                <item><title>과학 뉴스</title><link>{base}/article/2</link><pubDate>Tue, 06 Oct 2026 09:00:00 +0900</pubDate><category>과학</category></item>
                </channel></rss>'''
            elif self.path == '/news':
                body = '<main><a class="news" href="/article/1">First</a><a class="news" href="/article/2">Second</a></main>'
            elif self.path.startswith('/article/'):
                n = self.path.split('/')[-1].split('?')[0]
                body = f'''<html><head><meta property="article:published_time" content="2026-10-0{8-int(n)}T00:00:00Z"><meta property="article:section" content="{'우주' if n=='1' else '과학'}"></head>
                <body><h1>{'우주' if n=='1' else '과학'} 뉴스</h1><article><p>연구 결과 본문 {n} 버전 {state['version']}. 새로운 관측 결과를 발표했다.</p><script>ignore me</script></article></body></html>'''
            else:
                self.send_error(503)
                return
            self.send_response(200)
            self.send_header('Content-Type', 'text/xml; charset=utf-8' if self.path == '/feed' else 'text/html; charset=utf-8')
            self.end_headers()
            self.wfile.write(body.encode())

        def do_POST(self):
            payload = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
            state['ai_calls'] += 1
            state.setdefault('prompts', []).append(payload['messages'][0]['content'])
            assert payload['model'] == 'test-model'
            assert self.headers['Authorization'] == 'Bearer test-key'
            if state['fail_ai']:
                self.send_error(429)
                return
            text = '한국어 뉴스 요약입니다.'
            if 'trends' in payload['messages'][0]['content']:
                text = json.dumps({'trends': ['우주 연구 확대'], 'keywords': ['우주', '과학'], 'implications': ['관측 기술 발전']}, ensure_ascii=False)
            reply = {'choices': [{'message': {'content': text}, 'finish_reason': 'stop'}],
                     'usage': {'prompt_tokens': 100, 'completion_tokens': 20, 'total_tokens': 120}}
            if state.get('no_usage'):
                del reply['usage']
            if state.get('malformed_first') and state['ai_calls'] == 1:
                reply = {'choices': [None]}
            body = json.dumps(reply).encode()
            self.send_response(200)
            self.send_header('Content-Type', 'application/json')
            self.end_headers()
            self.wfile.write(body)

    server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    yield f'http://127.0.0.1:{server.server_port}', state
    server.shutdown()
    server.server_close()
    thread.join()


@pytest.fixture
def project(tmp_path, site):
    base, state = site
    cfg = {
        'database': 'news.db', 'output_dir': 'output', 'log_file': 'app.log',
        'duplicate_policy': 'skip', 'request_timeout': 2, 'request_delay': 0.01,
        'sources': {'test': {'rss_url': base + '/feed', 'crawl_url': base + '/news',
                            'link_selector': 'a.news', 'body_selector': 'article', 'category': '기타'}},
        'ai': {'base_url': base + '/v1', 'model': 'test-model', 'api_key': 'test-key',
               'api_key_env': 'TEST_NEWS_KEY', 'max_input_chars': 20000}
    }
    path = tmp_path / 'config.json'
    path.write_text(json.dumps(cfg), encoding='utf-8')
    return tmp_path, path, state


def run(project, *args, ok=True):
    assert (ROOT / 'main.py').exists(), 'The required CLI entry point has not been implemented'
    result = subprocess.run([sys.executable, str(ROOT / 'main.py'), '--config', str(project[1]), *args], cwd=ROOT, text=True, capture_output=True)
    if ok:
        assert result.returncode == 0, result.stdout + result.stderr
    else:
        assert result.returncode != 0
    return result


def rows(project, table):
    with sqlite3.connect(project[0] / 'news.db') as db:
        db.row_factory = sqlite3.Row
        return [dict(row) for row in db.execute(f'SELECT * FROM {table}')]


def seed(project):
    run(project, 'fetch', '--source', 'test', '--method', 'rss', '--limit', '2')
    run(project, 'clean')


def test_full_pipeline_and_persistent_outputs(project):
    seed(project)
    assert len(rows(project, 'raw_news')) == 2
    clean = rows(project, 'clean_news')
    assert len(clean) == 2
    assert clean[0]['published_at'] == '2026-10-07T00:00:00+00:00'
    assert 'utm_source' not in clean[0]['url']
    assert 'ignore me' not in clean[0]['body']
    run(project, 'summarize', '--unsummarized')
    assert all(row['summary'] for row in rows(project, 'clean_news'))
    assert project[2]['ai_calls'] == 2
    run(project, 'summarize', '--all')
    assert project[2]['ai_calls'] == 2  # Existing summaries must not be billed again.
    run(project, 'analyze', '--date-from', '2026-10-06', '--date-to', '2026-10-07')
    assert len(rows(project, 'analyses')) == 1
    result = run(project, 'report', '--date-from', '2026-10-06', '--date-to', '2026-10-07', '--top', '2')
    assert '우주 연구 확대' in result.stdout
    output = project[0] / 'output'
    assert (output / 'report.md').exists()
    assert len(list(output.rglob('*.png'))) == 2
    run(project, 'export', '--format', 'csv', '--status', 'summarized')
    run(project, 'export', '--format', 'excel', '--status', 'summarized')
    import csv
    with (output / 'news.csv').open(encoding='utf-8-sig') as f:
        assert len(list(csv.DictReader(f))) == 2
    from openpyxl import load_workbook
    wb = load_workbook(output / 'news.xlsx')
    assert wb.active.max_row == 3
    wb.close()
    run(project, 'export', '--format', 'jsonl', '--status', 'summarized')
    lines = (output / 'news.jsonl').read_text(encoding='utf-8').splitlines()
    assert len(lines) == 2
    first = json.loads(lines[0])
    assert first['title'] == '우주 뉴스' and first['summary']
    log = (project[0] / 'app.log').read_text(encoding='utf-8')
    assert '요약 토큰 합계: 호출 2회 / 입력 200 / 출력 40 / 합계 240' in log
    assert '분석 토큰 합계: 호출 1회' in log


def test_crawl_duplicates_and_upsert_invalidates_old_summary(project):
    seed(project)
    run(project, 'summarize', '--id', '1')
    run(project, 'fetch', '--source', 'test', '--method', 'crawl', '--limit', '2')
    run(project, 'clean')
    assert len(rows(project, 'clean_news')) == 2
    assert rows(project, 'clean_news')[0]['summary']
    project[2]['version'] = 2
    run(project, 'fetch', '--source', 'test', '--method', 'crawl', '--limit', '2')
    run(project, 'clean', '--duplicate-policy', 'upsert')
    clean = rows(project, 'clean_news')
    assert len(clean) == 2
    assert '버전 2' in clean[0]['body']
    assert clean[0]['summary'] is None
    before = rows(project, 'raw_news')
    run(project, 'clean')
    assert rows(project, 'raw_news') == before


def test_ai_failure_keeps_articles_retryable(project):
    seed(project)
    project[2]['fail_ai'] = True
    run(project, 'summarize', '--all', ok=False)
    assert project[2]['ai_calls'] == 2
    assert all(row['summary'] is None for row in rows(project, 'clean_news'))
    project[2]['fail_ai'] = False
    run(project, 'summarize', '--unsummarized')
    assert all(row['summary'] for row in rows(project, 'clean_news'))


def test_robots_denial_and_http_failure(project):
    project[2]['deny'] = True
    run(project, 'fetch', '--source', 'test', '--method', 'crawl', ok=False)
    assert not rows(project, 'raw_news')
    cfg = json.loads(project[1].read_text())
    cfg['sources']['test']['rss_url'] += '-unavailable'
    project[1].write_text(json.dumps(cfg))
    run(project, 'fetch', '--source', 'test', '--method', 'rss', ok=False)


def test_filters_empty_data_and_invalid_arguments(project):
    run(project, 'export', '--format', 'csv')
    run(project, 'report')
    run(project, 'analyze', ok=False)
    seed(project)
    run(project, 'summarize', '--all', '--category', '우주', '--date', '2026-10-07')
    assert sum(bool(row['summary']) for row in rows(project, 'clean_news')) == 1
    run(project, 'export', '--format', 'excel', '--category', '과학')
    run(project, 'fetch', '--limit', '-1', ok=False)
    run(project, 'analyze', '--date-from', '2026-11-01', '--date-to', '2026-10-01', ok=False)
    run(project, 'summarize', '--id', '999', ok=False)


def test_malformed_ai_response_does_not_abort_batch(project):
    seed(project)
    project[2]['malformed_first'] = True
    run(project, 'summarize', '--all', ok=False)
    assert project[2]['ai_calls'] == 2
    assert sum(bool(a['summary']) for a in rows(project, 'clean_news')) == 1


def test_txt_report_does_not_change_existing_markdown_charts(project):
    import re
    seed(project)
    run(project, 'report', '--category', '우주')
    output = project[0] / 'output'
    markdown = (output / 'report.md').read_text()
    links = re.findall(r'!\[[^\]]+\]\(([^)]+)\)', markdown)
    assert len(links) == 2
    images_before = [(output / name).read_bytes() for name in links]
    run(project, 'report', '--category', '과학', '--format', 'txt')
    assert '과학' in (output / 'report.txt').read_text()
    assert [(output / name).read_bytes() for name in links] == images_before


def test_limit_selects_remaining_work_and_unchanged_upsert_keeps_summary(project):
    seed(project)
    run(project, 'summarize', '--unsummarized', '--limit', '1')
    run(project, 'summarize', '--unsummarized', '--limit', '1')
    assert all(a['summary'] for a in rows(project, 'clean_news'))
    assert project[2]['ai_calls'] == 2
    run(project, 'fetch', '--source', 'test', '--method', 'rss', '--limit', '2')
    run(project, 'clean', '--duplicate-policy', 'upsert')
    run(project, 'summarize', '--all')
    assert project[2]['ai_calls'] == 2


def test_jsonl_export_filters_and_missing_usage(project):
    seed(project)
    project[2]['no_usage'] = True
    run(project, 'summarize', '--id', '1')
    run(project, 'export', '--format', 'jsonl', '--status', 'unsummarized')
    lines = (project[0] / 'output/news.jsonl').read_text(encoding='utf-8').splitlines()
    assert [json.loads(line)['id'] for line in lines] == [2]
    assert json.loads(lines[0])['summary'] is None
    assert 'usage 미제공 1회' in (project[0] / 'app.log').read_text(encoding='utf-8')


def test_prompt_override_from_config(project):
    seed(project)
    cfg = json.loads(project[1].read_text(encoding='utf-8'))
    cfg['ai']['summary_prompt'] = '실험 B: 한 문장으로 요약하세요.'
    project[1].write_text(json.dumps(cfg), encoding='utf-8')
    run(project, 'summarize', '--id', '1')
    assert project[2]['prompts'][-1].startswith('실험 B: 한 문장으로 요약하세요.')


def test_stale_analysis_and_export_filters(project):
    from openpyxl import load_workbook
    seed(project)
    run(project, 'summarize', '--all', '--category', '우주')
    run(project, 'analyze')
    project[2]['version'] = 2
    run(project, 'fetch', '--source', 'test', '--method', 'rss', '--limit', '2')
    run(project, 'clean', '--duplicate-policy', 'upsert')
    assert '현재 데이터와 분석 표본이 다릅니다' in run(project, 'report').stdout
    run(project, 'summarize', '--id', '2')
    run(project, 'export', '--format', 'excel', '--status', 'summarized', '--date', '2026-10-06', '--category', '과학')
    book = load_workbook(project[0] / 'output/news.xlsx')
    assert book.active.max_row == 2
    assert book.active['B2'].value == '과학 뉴스'
    assert book.active['D2'].value == '과학'
    book.close()
